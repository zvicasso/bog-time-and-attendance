# Time & Attendance Processor
# Author: Isheanesu Zvikaramba
# Copyright (c) 2026 Zvicasso Technologies (Pty) Ltd. All rights reserved.
"""Excel workbook generation (openpyxl).
"""

from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import branding

NAVY = branding.NAVY
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SECTION_FONT = Font(bold=True, size=11, color=NAVY)
MUTED_FONT = Font(italic=True, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)
ZEBRA = PatternFill("solid", fgColor="EEFAF9")

FMT = {
    "datetime": "dd/mm/yyyy hh:mm:ss",
    "date": "dd/mm/yyyy",
    "time": "hh:mm",
    "duration": "[h]:mm:ss",
    "number": "0.0",
    "int": "0",
    "text": "@",
}

TEXT_ALWAYS = {"Staff No", "Card Number(s)", "Record ID(s)", "Record ID", "Source Page(s)",
               "Duplicate Of"}

# (header, type, width hint)
RAW_COLS = [
    ("Record ID", "text", 15), ("Block ID", "int", 9), ("Source Page/Row", "int", 9),
    ("Cardholder Name", "text", 30), ("Staff No", "text", 13), ("Card Number(s)", "text", 24),
    ("Department (source)", "text", 24),
    ("Entry Clock Name", "text", 22), ("First Entry Date/Time", "datetime", 20),
    ("Exit Clock Name", "text", 22), ("Last Exit Date/Time", "datetime", 20),
    ("No Real Exit (placeholder)", "text", 12), ("Time Worked", "duration", 11),
    ("Time Worked (as printed)", "text", 12), ("Insufficient Data", "text", 10),
    ("Section Total (as printed)", "text", 12), ("Excluded", "text", 9),
    ("Exclusion Reason", "text", 34), ("Duplicate Of", "text", 15), ("Notes", "text", 34),
    ("Source Text", "text", 60),
]

LATE_COLS = ["Date", "Cardholder Name", "Staff No", "Card Number(s)", "Department", "On Roster",
             "First Entry", "Entry Clock", "Late Threshold", "Minutes Late", "Last Exit", "Exit Clock",
             "Also Early Departure", "Review Notes", "Source Page(s)", "Record ID(s)"]
EARLY_COLS = ["Date", "Cardholder Name", "Staff No", "Card Number(s)", "Department", "On Roster",
              "Last Exit", "Exit Clock", "Departure Threshold", "Minutes Early", "First Entry",
              "Entry Clock", "Also Late Arrival", "Review Notes", "Source Page(s)", "Record ID(s)"]
ONTIME_COLS = ["Date", "Cardholder Name", "Staff No", "Card Number(s)", "Department", "On Roster",
               "First Entry", "Entry Clock", "Last Exit", "Exit Clock", "Time on Site", "Review Notes",
               "Source Page(s)", "Record ID(s)"]
NOENTRY_COLS = ["Date", "Cardholder Name", "Staff No", "Card Number(s)", "Department", "On Roster",
                "Status", "Other Activity", "Source Page(s)", "Record ID(s)"]
EXC_COLS = ["Date", "Cardholder Name", "Staff No", "Card Number(s)", "Department", "On Roster",
            "Exception Type", "Detail", "First Entry", "Last Exit", "Source Page(s)", "Record ID(s)"]
ROSTER_COLS = [("Name (roster)", "text", 32), ("Staff No", "text", 12), ("Department", "text", 36),
               ("Status", "text", 58), ("Roster Reference", "text", 16)]
VALID_COLS = [("Severity", "text", 10), ("Category", "text", 20), ("Message", "text", 70),
              ("Page", "int", 7), ("Block ID", "int", 9), ("Source Text", "text", 60)]

TYPE_BY_NAME = {
    "Date": "date", "First Entry": "datetime", "Last Exit": "datetime",
    "Late Threshold": "time", "Departure Threshold": "time",
    "Minutes Late": "number", "Minutes Early": "number", "Time on Site": "duration",
}
WIDTH_BY_NAME = {"Date": 11, "Cardholder Name": 30, "Detail": 60, "Status": 44,
                 "Other Activity": 46, "Review Notes": 44, "Exception Type": 28, "Record ID(s)": 30}


def _style_header(ws: Worksheet, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _set_cell(cell, value, ctype: str, header: str):
    if header in TEXT_ALWAYS or ctype == "text":
        cell.value = "" if value is None else str(value)
        cell.number_format = "@"
        return
    if value is None or value == "":
        cell.value = None
        return
    if ctype == "time" and isinstance(value, time):
        cell.value = value
    elif ctype == "duration" and isinstance(value, timedelta):
        cell.value = value
    else:
        cell.value = value
    cell.number_format = FMT.get(ctype, "General")


def _print_setup(ws: Worksheet, landscape: bool = True, title_rows: str | None = None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.left.text = f"{branding.PRODUCT} - {branding.COMPANY_SHORT}".replace("&", "&&")
    ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = "&A - page &P of &N"


def _write_table(ws: Worksheet, columns: list[tuple[str, str, int]], rows: list[list],
                 empty_message: str, start_row: int = 1):
    ncols = len(columns)
    _print_setup(ws, title_rows=f"{start_row}:{start_row}")
    for i, (h, _, w) in enumerate(columns, start=1):
        ws.cell(row=start_row, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w
    _style_header(ws, start_row, ncols)
    if not rows:
        c = ws.cell(row=start_row + 1, column=1, value=empty_message)
        c.font = MUTED_FONT
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(ncols)}{start_row}"
    else:
        for r_i, values in enumerate(rows, start=start_row + 1):
            zebra = (r_i - start_row) % 2 == 0
            for c_i, ((h, t, _), v) in enumerate(zip(columns, values), start=1):
                cell = ws.cell(row=r_i, column=c_i)
                _set_cell(cell, v, t, h)
                if zebra:
                    cell.fill = ZEBRA
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(ncols)}{start_row + len(rows)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3 if ncols > 4 else 1)


def _category_sheet(wb: Workbook, title: str, names: list[str], rows: list[dict], roster: bool, msg: str):
    if not roster:
        names = [n for n in names if n != "On Roster"]
    if rows and not any(str(r.get("Card Number(s)") or "").strip() for r in rows):
        names = [n for n in names if n != "Card Number(s)"]
    if not any(r.get("Department") for r in rows) and not roster:
        names = [n for n in names if n != "Department"]
    cols = [(n, TYPE_BY_NAME.get(n, "text"), WIDTH_BY_NAME.get(n, max(12, min(24, len(n) + 4))))
            for n in names]
    ws = wb.create_sheet(title)
    _write_table(ws, cols, [[r.get(n) for n in names] for r in rows], msg)
    return ws


def _yn(b: bool) -> str:
    return "Yes" if b else "No"


def build_workbook(outcome) -> bytes:
    """Build the workbook from a pipeline.ProcessingOutcome and return its bytes."""
    a = outcome.analysis
    roster = a.roster_supplied
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _summary_sheet(ws, outcome)

    # Raw data (every parsed row, plus sections without rows, for traceability)
    raw_rows = []
    for b in outcome.parse.blocks:
        emp_excl = outcome.exclusions.get(b.block_id, "")
        if not b.sessions:
            raw_rows.append([f"B{b.block_id:05d}-R00", b.block_id, b.page, b.cardholder_name,
                             b.staff_no, b.card_numbers, b.department, "", None, "", None, "", None, "",
                             "", b.total_working_raw, _yn(bool(emp_excl)), emp_excl, "",
                             "Section has no attendance rows", ""])
        for s in b.sessions:
            exit_label = s.exit_clock or ("Report End Date/Time" if "Report End" in s.missing_exit_label else "")
            raw_rows.append([s.record_id, s.block_id, s.page, s.cardholder_name, s.staff_no,
                             s.card_numbers, b.department, s.entry_clock, s.entry_dt, exit_label,
                             s.exit_dt, _yn(s.exit_is_report_end), s.time_worked, s.time_worked_raw,
                             _yn(s.insufficient_data), b.total_working_raw, _yn(s.excluded),
                             s.exclusion_reason, s.duplicate_of, "; ".join(s.notes), s.raw_text])
    _write_table(wb.create_sheet("Raw Data"), RAW_COLS, raw_rows, "No records were extracted.")

    dates_txt = ", ".join(d.strftime("%d/%m/%Y") for d in a.dates)
    _category_sheet(wb, "Late Arrivals", LATE_COLS, a.late, roster,
                    f"No late arrivals for {dates_txt}.")
    _category_sheet(wb, "Early Departures", EARLY_COLS, a.early, roster,
                    f"No early departures for {dates_txt}.")
    _category_sheet(wb, "On Time", ONTIME_COLS, a.on_time, roster,
                    f"No employees arrived on time and left on time on {dates_txt}.")
    _category_sheet(wb, "No Recorded Entry", NOENTRY_COLS, a.no_entry, roster,
                    f"Every cardholder in the report has an entry on {dates_txt}.")
    _category_sheet(wb, "Attendance Exceptions", EXC_COLS, a.exceptions, roster,
                    f"No attendance exceptions for {dates_txt}.")

    if roster:
        rrows = [[e.name, e.staff_no, e.department,
                  "On roster, no section in attendance report (no badge activity, or not selected "
                  "when the report was run)", outcome.roster.ref(e)]
                 for e in sorted(outcome.roster_unmatched, key=lambda e: (e.department.lower(), e.name.lower()))]
        _write_table(wb.create_sheet("Roster Not In Report"), ROSTER_COLS, rrows,
                     "Every roster employee has a section in the attendance report.")

    vrows = [[i.severity, i.category, i.message, i.page, i.block_id, i.raw_text]
             for i in outcome.all_issues]
    _write_table(wb.create_sheet("Validation"), VALID_COLS, vrows,
                 "No parsing or validation problems were found.")

    wb.properties.title = "Time and Attendance Analysis"
    wb.properties.creator = f"{branding.AUTHOR}, {branding.COMPANY_SHORT}"
    wb.properties.description = f"Generated by {branding.PRODUCT} - developed by {branding.AUTHOR}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary_sheet(ws: Worksheet, o):
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 18
    r = 1
    if branding.LOGO_HORIZONTAL.exists():
        from openpyxl.drawing.image import Image as XLImage

        logo = XLImage(str(branding.LOGO_HORIZONTAL))
        logo.width, logo.height = 200, int(200 * logo.height / logo.width)
        ws.add_image(logo, "A1")
        for i in range(1, 5):
            ws.row_dimensions[i].height = 19
        r = 6
    ws.cell(r, 1, "Time and Attendance Analysis").font = TITLE_FONT
    r += 1
    ws.cell(r, 1, o.meta.site_name or "").font = MUTED_FONT
    r += 1
    ws.cell(r, 1, f"Prepared with {branding.PRODUCT} · developed by {branding.AUTHOR}, "
                  f"{branding.COMPANY_SHORT}").font = Font(size=8, color="808080")
    r += 2

    def section(title):
        nonlocal r
        ws.cell(r, 1, title).font = SECTION_FONT
        ws.cell(r, 1).border = Border(bottom=Side(style="medium", color=NAVY))
        ws.cell(r, 2).border = Border(bottom=Side(style="medium", color=NAVY))
        r += 1

    def kv(k, v, fmt=None):
        nonlocal r
        ws.cell(r, 1, k).alignment = Alignment(vertical="top")
        c = ws.cell(r, 2, v)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if fmt:
            c.number_format = fmt
        r += 1

    s = o.summary
    section("Processing information")
    for k in ("Source file", "Source format", "Staff ID field", "Site", "Report period", "Report generated",
              "Pages / rows read", "Reporting date(s)", "Late arrival threshold", "Early departure threshold",
              "Exclusion keywords", "Visitor doors", "Employee roster", "Processed at", "Processing status"):
        kv(k, s.get(k, ""))
    r += 1
    section("Record counts")
    for k in ("Cardholder sections in report", "Total records extracted", "Unique cardholders",
              "Excluded cardholders", "Excluded records", "Employees analysed", "Duplicate records ignored",
              "On time (arrived and left on time)", "Late arrivals", "Early departures", "No recorded entry", "Attendance exceptions (rows)",
              "Employees with exceptions", "Roster employees not in attendance report",
              "Validation warnings / errors"):
        kv(k, s.get(k, 0), "0")
    r += 1
    section("Exceptions by type")
    if o.exception_counts:
        for k, v in sorted(o.exception_counts.items(), key=lambda kv_: -kv_[1]):
            kv(k, v, "0")
    else:
        kv("None", 0, "0")
    r += 1
    section("How to read this workbook")
    for line in o.interpretation_notes:
        c = ws.cell(r, 1, "•  " + line)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    section("Warnings")
    if o.warnings:
        for w in o.warnings[:60]:
            c = ws.cell(r, 1, w)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        if len(o.warnings) > 60:
            ws.cell(r, 1, f"... {len(o.warnings) - 60} more - see the Validation sheet.").font = MUTED_FONT
    else:
        ws.cell(r, 1, "None").font = MUTED_FONT
    ws.sheet_view.showGridLines = False
    _print_setup(ws, landscape=False)
