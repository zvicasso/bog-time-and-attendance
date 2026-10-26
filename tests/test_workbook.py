# Time & Attendance Processor
# Author: Isheanesu Zvikaramba
# Copyright (c) 2026 Zvicasso Technologies (Pty) Ltd. All rights reserved.
import io
from datetime import date

import openpyxl
import pytest

from tap.pipeline import process_pdf

EXPECTED = ["Summary", "Raw Data", "Late Arrivals", "Early Departures", "On Time", "No Recorded Entry",
            "Attendance Exceptions", "Validation"]


@pytest.fixture(scope="module")
def wb(synthetic_pdf):
    from tap.config import Settings
    o = process_pdf(synthetic_pdf, Settings(), dates=[date(2026, 10, 21)])
    return openpyxl.load_workbook(io.BytesIO(o.workbook))


def headers(ws):
    return [c.value for c in ws[1]]


def test_sheet_order(wb):
    assert wb.sheetnames == EXPECTED


def test_raw_data_traceability_and_text_ids(wb):
    ws = wb["Raw Data"]
    h = headers(ws)
    assert h[:3] == ["Record ID", "Block ID", "Source Page/Row"]
    staff = h.index("Staff No") + 1
    first = ws.cell(row=2, column=staff)
    assert first.value == "000123" and first.number_format == "@"
    cards = ws.cell(row=2, column=h.index("Card Number(s)") + 1)
    assert isinstance(cards.value, str)
    # 36 rows + 1 placeholder for the section without rows
    assert ws.max_row == 1 + 37
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref.startswith("A1:")


def test_datetime_formats(wb):
    ws = wb["Late Arrivals"]
    h = headers(ws)
    c = ws.cell(row=2, column=h.index("First Entry") + 1)
    assert c.number_format == "dd/mm/yyyy hh:mm:ss"
    assert ws.cell(row=2, column=1).number_format == "dd/mm/yyyy"
    raw = wb["Raw Data"]
    tw = raw.cell(row=2, column=headers(raw).index("Time Worked") + 1)
    assert tw.number_format == "[h]:mm:ss"


def test_no_roster_columns_without_roster(wb):
    assert "On Roster" not in headers(wb["No Recorded Entry"])


def test_empty_category_still_has_headers(synthetic_pdf, settings):
    from datetime import time
    o = process_pdf(synthetic_pdf, settings.with_overrides(late_threshold=time(23, 0)),
                    dates=[date(2026, 10, 21)])
    w = openpyxl.load_workbook(io.BytesIO(o.workbook))
    ws = w["Late Arrivals"]
    assert headers(ws)[0] == "Date"
    assert "No late arrivals" in ws.cell(row=2, column=1).value


def test_summary_counts(wb):
    ws = wb["Summary"]
    vals = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0]}
    assert vals["Late arrivals"] == 4
    assert vals["Early departures"] == 2
    assert vals["No recorded entry"] == 2
    assert vals["Excluded cardholders"] == 4
